#!/usr/bin/env python3
# pptx_extract.py — extract text from a .pptx via python-pptx.
#   python3 pptx_extract.py <in.pptx>  → slide text on stdout
import sys
from pptx import Presentation

def main(path):
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        print(f'--- slide {i} ---')
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = ''.join(run.text for run in para.runs).strip()
                    if line:
                        print(line)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [' '.join(p.text for p in c.text_frame.paragraphs).strip() for c in row.cells]
                    print(' | '.join(cells))

if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.stderr.write('usage: pptx_extract.py <in.pptx>\n')
        sys.exit(2)
    main(sys.argv[1])
