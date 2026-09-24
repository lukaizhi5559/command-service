#!/usr/bin/env python3
# docx_ops.py — deterministic .docx read/edit helper for edit.agent.
#
#   python3 docx_ops.py extract <in.docx>          → numbered paragraphs + tables on stdout
#   python3 docx_ops.py apply <in.docx> <out.docx> → reads {"ops":[...]} from stdin,
#                                                    writes edited copy to out.docx,
#                                                    prints {"applied":N,"missed":M}
#
# Ops (docx):
#   {"find": "<verbatim text>", "replace": "<new text>", "all": false}
#   {"append": "<new paragraph text>"}
#
# Find/replace works at paragraph level: the first paragraph whose .text
# contains `find` gets its runs rebuilt with the replacement. Inline formatting
# inside an edited paragraph is flattened (single run) — acceptable because
# output is always a reviewable draft copy; originals are never written.
# With "all": true, every matching paragraph/cell is edited.

import sys
import os
import json
import shutil

import docx


def _set_para_text(p, text):
    # Rebuild the paragraph as a single run — flattens inline formatting
    # (bold/italic spans) inside the edited paragraph only.
    for r in p.runs:
        r.text = ''
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def _iter_paragraphs(d):
    for p in d.paragraphs:
        yield p
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def extract(path):
    d = docx.Document(path)
    for i, p in enumerate(d.paragraphs):
        if p.text.strip():
            print(f"[{i}] {p.text}")
    for ti, tbl in enumerate(d.tables):
        for ri, row in enumerate(tbl.rows):
            cells = " | ".join(c.text for c in row.cells)
            if cells.strip():
                print(f"[T{ti}R{ri}] {cells}")


def apply_ops(src, dst):
    payload = json.load(sys.stdin)
    ops = payload.get('ops', [])[:40]
    # Converted drafts (rtf/doc) live at workPath — src==dst means apply in place.
    if os.path.realpath(src) != os.path.realpath(dst):
        shutil.copyfile(src, dst)
    d = docx.Document(dst)
    applied = 0
    missed = 0
    for op in ops:
        if not isinstance(op, dict):
            missed += 1
            continue
        if 'append' in op:
            d.add_paragraph(str(op['append']))
            applied += 1
            continue
        find = str(op.get('find', ''))
        repl = str(op.get('replace', ''))
        if not find:
            missed += 1
            continue
        hit = False
        for p in _iter_paragraphs(d):
            if find in p.text:
                _set_para_text(p, p.text.replace(find, repl))
                applied += 1
                hit = True
                if not op.get('all'):
                    break
        if not hit:
            missed += 1
    d.save(dst)
    print(json.dumps({'applied': applied, 'missed': missed}))


if __name__ == '__main__':
    if len(sys.argv) < 3:
        sys.stderr.write('usage: docx_ops.py extract <in.docx> | apply <in.docx> <out.docx>\n')
        sys.exit(2)
    cmd = sys.argv[1]
    if cmd == 'extract':
        extract(sys.argv[2])
    elif cmd == 'apply' and len(sys.argv) >= 4:
        apply_ops(sys.argv[2], sys.argv[3])
    else:
        sys.stderr.write('usage: docx_ops.py extract <in.docx> | apply <in.docx> <out.docx>\n')
        sys.exit(2)
